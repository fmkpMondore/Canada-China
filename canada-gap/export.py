"""
Export analysis results to formatted Excel or CSV.
"""

import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter

from config import OUTPUT_DIR


# Colours
HEADER_FILL = PatternFill("solid", fgColor="1F3864")   # dark navy
ALT_FILL    = PatternFill("solid", fgColor="DCE6F1")   # light blue
TOP5_FILL   = PatternFill("solid", fgColor="FFF2CC")   # light yellow
TOP1_FILL   = PatternFill("solid", fgColor="FFD966")   # gold

HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
BODY_FONT   = Font(name="Calibri", size=10)

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _apply_header(ws, row: int, columns: list[str]):
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _apply_body_row(ws, row: int, values: list, is_alt: bool, is_top5: bool, is_top1: bool):
    fill = TOP1_FILL if is_top1 else (TOP5_FILL if is_top5 else (ALT_FILL if is_alt else PatternFill()))
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col_idx, value=value)
        cell.font = BODY_FONT
        cell.fill = fill
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center")


def _autofit_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                cell_len = len(str(cell.value)) if cell.value is not None else 0
                if cell_len > max_len:
                    max_len = cell_len
            except Exception:
                pass
        adjusted = min(max_len + 4, 60)
        ws.column_dimensions[col_letter].width = adjusted


def _build_summary_sheet(ws, df: pd.DataFrame, title: str):
    """Write a ranked summary sheet."""
    ws.title = title[:31]  # Excel sheet name limit

    columns = [
        "Rank", "HS Code", "Description",
        "World (USD M)", "China (USD M)", "Gap (USD M)",
        "China Share %", "Opportunity Score (B USD)",
    ]
    _apply_header(ws, 1, columns)
    ws.freeze_panes = "A2"

    for i, (rank, row) in enumerate(df.iterrows()):
        excel_row = i + 2
        values = [
            rank,
            row["hs_code"],
            row["description"],
            round(row["world_usd"] / 1e6, 2),
            round(row["china_usd"] / 1e6, 2),
            round(row["gap_usd"] / 1e6, 2),
            round(row["china_share_pct"], 2),
            round(row["opportunity_score"] / 1e9, 3),
        ]
        _apply_body_row(
            ws, excel_row, values,
            is_alt=(i % 2 == 1),
            is_top5=(rank <= 5),
            is_top1=(rank == 1),
        )

    # Number formats
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = '#,##0.00'
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=8):
        for cell in row:
            cell.number_format = '0.00'

    _autofit_columns(ws)


def _build_yearly_sheet(ws, df: pd.DataFrame):
    """Write yearly detail sheet."""
    ws.title = "By Year"

    columns = ["Year", "HS Code", "Description", "World (USD M)", "China (USD M)", "Gap (USD M)", "China Share %"]
    _apply_header(ws, 1, columns)
    ws.freeze_panes = "A2"

    for i, (_, row) in enumerate(df.iterrows()):
        excel_row = i + 2
        values = [
            int(row["year"]) if pd.notna(row.get("year")) else "",
            row["hs_code"],
            row.get("description", ""),
            round(row["world_usd"] / 1e6, 2),
            round(row["china_usd"] / 1e6, 2),
            round(row["gap_usd"] / 1e6, 2),
            round(row["china_share_pct"], 2),
        ]
        _apply_body_row(ws, excel_row, values, is_alt=(i % 2 == 1), is_top5=False, is_top1=False)

    _autofit_columns(ws)


def export_holistic(
    summary_df: pd.DataFrame,
    yearly_df: pd.DataFrame,
    filename: str = "canada_import_gap_holistic.xlsx",
) -> str:
    """Export holistic analysis to Excel with two sheets."""
    path = os.path.join(OUTPUT_DIR, filename)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        # Placeholder sheets (will be overwritten)
        summary_df.reset_index().to_excel(writer, sheet_name="Ranking", index=False)
        yearly_df.to_excel(writer, sheet_name="By Year", index=False)

    # Re-open and apply formatting
    wb = load_workbook(path)
    ws_summary = wb["Ranking"]
    ws_yearly = wb["By Year"]

    # Clear and rebuild
    ws_summary.delete_rows(1, ws_summary.max_row)
    ws_yearly.delete_rows(1, ws_yearly.max_row)

    _build_summary_sheet(ws_summary, summary_df, "Ranking")
    _build_yearly_sheet(ws_yearly, yearly_df)

    wb.save(path)
    return path


def export_deepdive(
    summary_df: pd.DataFrame,
    chapter: str,
    filename: str = None,
) -> str:
    """Export deep dive analysis to Excel."""
    if filename is None:
        filename = f"canada_import_gap_hs{chapter}_deepdive.xlsx"
    path = os.path.join(OUTPUT_DIR, filename)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary_df.reset_index().to_excel(writer, sheet_name="DeepDive", index=False)

    wb = load_workbook(path)
    ws = wb["DeepDive"]
    ws.delete_rows(1, ws.max_row)
    _build_summary_sheet(ws, summary_df, f"HS{chapter} Deep Dive")
    wb.save(path)
    return path


def export_csv(df: pd.DataFrame, filename: str) -> str:
    """Export DataFrame to CSV."""
    path = os.path.join(OUTPUT_DIR, filename)
    df.reset_index().to_csv(path, index=False)
    return path
